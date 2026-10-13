import sys
import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── Color Palette (PayGate Fintech Enterprise) ──────────────────────────────────
COLOR_BG_DARK       = RGBColor(15, 23, 42)      # Deep Navy #0F172A
COLOR_BG_LIGHT      = RGBColor(248, 250, 252)   # Light Slate Gray #F8FAFC
COLOR_WHITE         = RGBColor(255, 255, 255)
COLOR_NAVY_DARK     = RGBColor(17, 24, 39)      # Charcoal Navy #111827
COLOR_NAVY_PRIMARY  = RGBColor(30, 58, 138)     # Deep Blue #1E3A8A
COLOR_BLUE_ACCENT   = RGBColor(37, 99, 235)     # Royal Blue #2563EB
COLOR_TEAL_ACCENT   = RGBColor(13, 148, 136)    # Dark Teal #0D9488
COLOR_SLATE_BORDER  = RGBColor(226, 232, 240)   # Slate Border #E2E8F0
COLOR_SLATE_TEXT    = RGBColor(71, 85, 105)     # Slate Muted #475569
COLOR_SLATE_CARD    = RGBColor(255, 255, 255)   # Clean White Card
COLOR_CARD_ALT      = RGBColor(241, 245, 249)   # Slate-100 #F1F5F9

# Severity Colors
COLOR_SEV_LOW       = RGBColor(100, 116, 139)   # Slate #64748B
COLOR_SEV_MED       = RGBColor(217, 119, 6)     # Amber #D97706
COLOR_SEV_HIGH      = RGBColor(234, 88, 12)     # Orange #EA580C
COLOR_SEV_CRIT      = RGBColor(225, 29, 72)     # Crimson Rose #E11D48

FONT_HEADING = 'Arial'
FONT_BODY = 'Arial'

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_header_footer(slide, title_text, category_text="MODULE 06 — ADMIN OPERATIONS, MONITORING & ERROR AUDIT CONSOLE", slide_num=1):
        # Header category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
        p_c = tf_cat.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.name = FONT_BODY
        p_c.font.size = Pt(10)
        p_c.font.bold = True
        p_c.font.color.rgb = COLOR_BLUE_ACCENT

        # Slide Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.65))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_NAVY_DARK

        # Top Accent Divider
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.42), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_SLATE_BORDER
        line.line.color.rgb = COLOR_SLATE_BORDER

        # Footer
        footer_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.015))
        footer_line.fill.solid()
        footer_line.fill.fore_color.rgb = COLOR_SLATE_BORDER
        footer_line.line.color.rgb = COLOR_SLATE_BORDER

        f_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(8.0), Inches(0.3))
        tf_f = f_box.text_frame
        tf_f.margin_left = tf_f.margin_right = tf_f.margin_top = tf_f.margin_bottom = 0
        p_f = tf_f.paragraphs[0]
        p_f.text = "PayGate · Academic Project Evaluation · Module 6: Admin Operations & Error Audit"
        p_f.font.name = FONT_BODY
        p_f.font.size = Pt(9)
        p_f.font.color.rgb = COLOR_SLATE_TEXT

        p_num_box = slide.shapes.add_textbox(Inches(11.0), Inches(7.05), Inches(1.533), Inches(0.3))
        tf_num = p_num_box.text_frame
        tf_num.margin_left = tf_num.margin_right = tf_num.margin_top = tf_num.margin_bottom = 0
        p_num = tf_num.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.text = f"Slide {slide_num} of 7"
        p_num.font.name = FONT_BODY
        p_num.font.size = Pt(9)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_SLATE_TEXT

    # ════════════════════════════════════════════════════════════════════════════
    # SLIDE 1: Admin Console Overview
    # ════════════════════════════════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank_layout)
    # Background
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_BG_LIGHT
    bg1.line.color.rgb = COLOR_BG_LIGHT

    add_header_footer(s1, "Operational Visibility & System Governance", slide_num=1)

    # 3 Structured Pillar Cards
    pillars = [
        ("Why Administrators Require Visibility", 
         "Mission-Critical Oversight",
         [
             "PayGate processes non-reversible fund movements requiring continuous runtime scrutiny.",
             "Prevents silent transaction drop-offs where network or bank partner timeouts leave transfers unconfirmed.",
             "Protects customer trust by guaranteeing rapid reconciliation of trapped balances."
         ],
         COLOR_NAVY_PRIMARY),
        ("Core Administrative Responsibilities",
         "Operational Control Scope",
         [
             "Platform Telemetry: Continuous monitoring of gateway uptime and MongoDB database connectivity.",
             "Failure Interception: Real-time triage of high-severity exceptions before customer escalation.",
             "Administrative Kill-Switch: Immediate security freeze enforcement on suspicious or compromised wallets."
         ],
         COLOR_BLUE_ACCENT),
        ("Zero-Trust Architecture Role",
         "Strict Role-Based Boundary",
         [
             "Enforces strict RBAC separation: Customers are physically restricted from administrative telemetry routes.",
             "Synchronous route interception via auth and adminOnly middleware prevents unauthorized access.",
             "Every administrative intervention is captured in immutable audit logs with resolvedBy tracking."
         ],
         COLOR_TEAL_ACCENT),
    ]

    for i, (title, tag, points, accent_col) in enumerate(pillars):
        left = Inches(0.8 + i * 4.0)
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(3.733), Inches(4.9))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = COLOR_SLATE_BORDER
        card.line.width = Pt(1.5)

        # Accent top bar
        bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.8), Inches(3.733), Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent_col
        bar.line.color.rgb = accent_col

        tb = s1.shapes.add_textbox(left + Inches(0.3), Inches(2.15), Inches(3.133), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = tag.upper()
        p0.font.size = Pt(9.5)
        p0.font.bold = True
        p0.font.color.rgb = accent_col

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_NAVY_DARK
        p1.space_after = Pt(14)

        for pt in points:
            p = tf.add_paragraph()
            p.text = f"•  {pt}"
            p.font.size = Pt(10.5)
            p.font.color.rgb = COLOR_SLATE_TEXT
            p.space_after = Pt(10)
            p.line_spacing = 1.15

    # Speaker Notes Slide 1
    s1.notes_slide.notes_text_frame.text = (
        "Good morning respected evaluators. As part of our team's PayGate presentation, "
        "I am presenting Module 6: The Admin Operations, Monitoring & Error Audit Console.\n\n"
        "In PayGate, operational visibility is not optional—it is a regulatory and technical necessity. "
        "Without centralized oversight, transaction failures caused by core banking timeouts or frozen accounts "
        "can be silently dropped, leaving customer funds in an ambiguous state. "
        "The PayGate Admin Console acts as the operational nerve center, providing end-to-end telemetry, "
        "real-time error triage, and immediate kill-switch intervention under a zero-trust RBAC architecture."
    )

    # ════════════════════════════════════════════════════════════════════════════
    # SLIDE 2: Admin Dashboard (KPIs, Telemetry & Form Layout)
    # ════════════════════════════════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank_layout)
    bg2 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = COLOR_BG_LIGHT
    bg2.line.color.rgb = COLOR_BG_LIGHT

    add_header_footer(s2, "Admin Dashboard Architecture & Operational Console", slide_num=2)

    # Left Column: Structured Form Screenshot representation from Project Record
    left_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(6.0), Inches(4.9))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = COLOR_WHITE
    left_card.line.color.rgb = COLOR_SLATE_BORDER
    left_card.line.width = Pt(1.5)

    tb_snap = s2.shapes.add_textbox(Inches(1.05), Inches(1.95), Inches(5.5), Inches(4.6))
    tf_snap = tb_snap.text_frame
    tf_snap.word_wrap = True
    tf_snap.margin_left = tf_snap.margin_right = tf_snap.margin_top = tf_snap.margin_bottom = 0

    p_s0 = tf_snap.paragraphs[0]
    p_s0.text = "ADMIN CONSOLE UI SPECIFICATION (FORM 5)"
    p_s0.font.size = Pt(10)
    p_s0.font.bold = True
    p_s0.font.color.rgb = COLOR_BLUE_ACCENT

    p_s1 = tf_snap.add_paragraph()
    p_s1.text = "Documented Console Wireframe & Telemetry Layout"
    p_s1.font.size = Pt(13)
    p_s1.font.bold = True
    p_s1.font.color.rgb = COLOR_NAVY_DARK
    p_s1.space_after = Pt(10)

    # Code/ASCII container
    snap_box = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.05), Inches(2.7), Inches(5.5), Inches(3.75))
    snap_box.fill.solid()
    snap_box.fill.fore_color.rgb = COLOR_CARD_ALT
    snap_box.line.color.rgb = COLOR_SLATE_BORDER

    tb_code = s2.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(5.2), Inches(3.55))
    tf_code = tb_code.text_frame
    tf_code.word_wrap = True
    tf_code.margin_left = tf_code.margin_right = tf_code.margin_top = tf_code.margin_bottom = 0

    code_lines = [
        "+------------------------------------------------------------+",
        "| PAYGATE ADMIN OPERATIONS CONSOLE                           |",
        "| Telemetry: Live (Active) | Success Rate: 98.4% | Errors: 1 |",
        "+------------------------------------------------------------+",
        "| HIGH SEVERITY ERROR TELEMETRY                              |",
        "| +--------------------------------------------------------+ |",
        "| | Incident: ERR-9281 | Severity: HIGH                    | |",
        "| | Code: INSUFFICIENT_FUNDS | Affected: TXN-171829011     | |",
        "| | Message: Requested ₹50,000 exceeding balance ₹10,000   | |",
        "| | Status: [ Unresolved ]              [ Mark Resolved ]  | |",
        "| +--------------------------------------------------------+ |",
        "+------------------------------------------------------------+"
    ]
    for idx, cl in enumerate(code_lines):
        p_cl = tf_code.paragraphs[0] if idx == 0 else tf_code.add_paragraph()
        p_cl.text = cl
        p_cl.font.name = 'Consolas'
        p_cl.font.size = Pt(8.5)
        p_cl.font.color.rgb = COLOR_NAVY_DARK

    # Right Column: 4 Key Dashboard Monitoring Components
    kpi_blocks = [
        ("1. Transaction Volume & Settled Value", 
         "Tracks total count (completed, failed, pending) and settled volume in INR (₹) via GET /api/admin/stats."),
        ("2. Real-Time Success & Failure Rates",
         "Computes dynamic success percentages (e.g. 98.4%) to immediately alert operators to partner banking API latency or degradation."),
        ("3. In-Flight Pending Monitor",
         "Isolates unresolved payments caught in gateway clearing transitions with TTL tracking (Time-To-Live up to 30 mins)."),
        ("4. System Telemetry & Heartbeat",
         "Continuous 10-second polling via /health monitoring MongoDB status, DB ping latency, memory heap usage, and active uptime."),
    ]

    for j, (kpi_title, kpi_desc) in enumerate(kpi_blocks):
        top_pos = Inches(1.8 + j * 1.25)
        k_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), top_pos, Inches(5.433), Inches(1.12))
        k_card.fill.solid()
        k_card.fill.fore_color.rgb = COLOR_WHITE
        k_card.line.color.rgb = COLOR_SLATE_BORDER
        k_card.line.width = Pt(1.5)

        k_side = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.1), top_pos, Inches(0.12), Inches(1.12))
        k_side.fill.solid()
        k_side.fill.fore_color.rgb = COLOR_TEAL_ACCENT if j % 2 == 0 else COLOR_BLUE_ACCENT
        k_side.line.color.rgb = k_side.fill.fore_color.rgb

        tb_k = s2.shapes.add_textbox(Inches(7.38), top_pos + Inches(0.12), Inches(5.0), Inches(0.9))
        tf_k = tb_k.text_frame
        tf_k.word_wrap = True
        tf_k.margin_left = tf_k.margin_right = tf_k.margin_top = tf_k.margin_bottom = 0

        p_kt = tf_k.paragraphs[0]
        p_kt.text = kpi_title
        p_kt.font.size = Pt(11.5)
        p_kt.font.bold = True
        p_kt.font.color.rgb = COLOR_NAVY_DARK
        p_kt.space_after = Pt(3)

        p_kd = tf_k.add_paragraph()
        p_kd.text = kpi_desc
        p_kd.font.size = Pt(9.5)
        p_kd.font.color.rgb = COLOR_SLATE_TEXT
        p_kd.line_spacing = 1.15

    # Speaker Notes Slide 2
    s2.notes_slide.notes_text_frame.text = (
        "Slide 2 showcases the PayGate Admin Dashboard. On the left is the exact Form 5 layout from our project documentation. "
        "The interface aggregates four vital telemetry streams:\n"
        "First, system-wide transaction volume and settled value in INR from GET /api/admin/stats.\n"
        "Second, real-time success rate analytics, allowing operators to detect anomalies before users report them.\n"
        "Third, an active pending transactions counter showing in-flight transfers awaiting bank acknowledgments.\n"
        "And fourth, real-time system health telemetry polled every 10 seconds via /health, monitoring MongoDB ping latency, heap memory, and process uptime."
    )

    # ════════════════════════════════════════════════════════════════════════════
    # SLIDE 3: Transaction Monitoring & Ledger Inspection
    # ════════════════════════════════════════════════════════════════════════════
    s3 = prs.slides.add_slide(blank_layout)
    bg3 = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg3.fill.solid()
    bg3.fill.fore_color.rgb = COLOR_BG_LIGHT
    bg3.line.color.rgb = COLOR_BG_LIGHT

    add_header_footer(s3, "Comprehensive Transaction Monitoring & Audit Fields", slide_num=3)

    # Top summary card
    top_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(11.733), Inches(1.0))
    top_box.fill.solid()
    top_box.fill.fore_color.rgb = COLOR_WHITE
    top_box.line.color.rgb = COLOR_SLATE_BORDER
    top_box.line.width = Pt(1.5)

    tb_top = s3.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(11.1), Inches(0.7))
    tf_top = tb_top.text_frame
    tf_top.word_wrap = True
    tf_top.margin_left = tf_top.margin_right = tf_top.margin_top = tf_top.margin_bottom = 0
    p_t1 = tf_top.paragraphs[0]
    p_t1.text = "THE IMMUTABLE AUDIT LEDGER (TRANSACTIONS COLLECTION)"
    p_t1.font.size = Pt(10)
    p_t1.font.bold = True
    p_t1.font.color.rgb = COLOR_BLUE_ACCENT
    p_t2 = tf_top.add_paragraph()
    p_t2.text = "Administrators inspect transactions across all lifecycle states. Every record maintains cryptographic uniqueness and audit traceability."
    p_t2.font.size = Pt(11)
    p_t2.font.color.rgb = COLOR_NAVY_DARK

    # 4 Detailed Inspection Dimension Cards
    dimensions = [
        ("1. Unique Reference Code",
         "Format: TXN-<Timestamp>-<UUID>",
         [
             "Cryptographically unique reference code generated pre-save.",
             "Enables instantaneous single-record indexing and multi-party reconciliation.",
             "Used across customer support and banking partner dispute tickets."
         ]),
        ("2. State & Amount Auditing",
         "Status & Type Schema Enums",
         [
             "Tracks exact settlement sum in INR with two decimal precision.",
             "Enforces lifecycle status transitions: pending → success, failure, or cancelled.",
             "Captures transfer type: transfer, reversal, refund, topup, or enquiry."
         ]),
        ("3. Channel & Routing Method",
         "Documented Payment Methods",
         [
             "Monitors payment rails: WALLET, BANK_TRANSFER, CARD, UPI, NEFT, RTGS, IMPS.",
             "Verifies whether debits originated from Wallet or Linked Bank Accounts.",
             "Links senderWalletId to beneficiary recipientId with populated user details."
         ]),
        ("4. Failure Root-Cause Reason",
         "Diagnostic Field & Override Audit",
         [
             "Stores exact error rationale if execution aborts (e.g. INSUFFICIENT_FUNDS).",
             "Links directly to the ErrorLog document via transactionId.",
             "Records resolving administrator ID (resolvedBy) upon manual intervention."
         ]),
    ]

    for d_idx, (d_title, d_sub, d_items) in enumerate(dimensions):
        col_left = Inches(0.8 + d_idx * 3.0)
        c_shape = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col_left, Inches(2.95), Inches(2.78), Inches(3.75))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = COLOR_WHITE
        c_shape.line.color.rgb = COLOR_SLATE_BORDER
        c_shape.line.width = Pt(1.5)

        tb_d = s3.shapes.add_textbox(col_left + Inches(0.2), Inches(3.1), Inches(2.38), Inches(3.45))
        tf_d = tb_d.text_frame
        tf_d.word_wrap = True
        tf_d.margin_left = tf_d.margin_right = tf_d.margin_top = tf_d.margin_bottom = 0

        p_dt = tf_d.paragraphs[0]
        p_dt.text = d_title
        p_dt.font.size = Pt(11.5)
        p_dt.font.bold = True
        p_dt.font.color.rgb = COLOR_NAVY_DARK

        p_ds = tf_d.add_paragraph()
        p_ds.text = d_sub
        p_ds.font.size = Pt(9.5)
        p_ds.font.bold = True
        p_ds.font.color.rgb = COLOR_BLUE_ACCENT
        p_ds.space_after = Pt(10)

        for di in d_items:
            p_di = tf_d.add_paragraph()
            p_di.text = f"• {di}"
            p_di.font.size = Pt(9.5)
            p_di.font.color.rgb = COLOR_SLATE_TEXT
            p_di.space_after = Pt(8)
            p_di.line_spacing = 1.15

    # Speaker Notes Slide 3
    s3.notes_slide.notes_text_frame.text = (
        "In Slide 3, we examine how administrators inspect the transaction ledger in PayGate. "
        "The system stores all operations in MongoDB's transactions collection with high indexing. "
        "Administrators can filter by unique reference code generated with timestamps and UUIDs, "
        "inspect amounts and terminal states, trace whether the transfer occurred via Wallet or IMPS, "
        "and immediately inspect the failure reason. This complete visibility ensures zero ambiguity "
        "when auditing transaction disputes."
    )

    # ════════════════════════════════════════════════════════════════════════════
    # SLIDE 4: Error Telemetry Lifecycle Flow
    # ════════════════════════════════════════════════════════════════════════════
    s4 = prs.slides.add_slide(blank_layout)
    bg4 = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg4.fill.solid()
    bg4.fill.fore_color.rgb = COLOR_BG_LIGHT
    bg4.line.color.rgb = COLOR_BG_LIGHT

    add_header_footer(s4, "Error Telemetry Pipeline & Database Schema", slide_num=4)

    # Flow Steps
    flow_steps = [
        ("1. Transaction Failure", "Payment debit fails due to business rule or partner error", COLOR_SEV_CRIT),
        ("2. Error Classification", "failureReason parsed and categorized automatically", COLOR_SEV_HIGH),
        ("3. ErrorLog Creation", "Mongoose post-save hook creates audit document", COLOR_NAVY_PRIMARY),
        ("4. Severity Rating", "Assigned: LOW, MEDIUM, HIGH, or CRITICAL", COLOR_SEV_MED),
        ("5. Admin Review", "Console alerts on unresolved error telemetry via /admin/errors", COLOR_BLUE_ACCENT),
        ("6. Resolution", "Operator investigates and resolves via PATCH /admin/errors/:id/resolve", COLOR_TEAL_ACCENT),
    ]

    for f_idx, (f_title, f_sub, f_col) in enumerate(flow_steps):
        f_left = Inches(0.8 + f_idx * 1.98)
        f_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, f_left, Inches(1.8), Inches(1.8), Inches(1.7))
        f_box.fill.solid()
        f_box.fill.fore_color.rgb = COLOR_WHITE
        f_box.line.color.rgb = f_col
        f_box.line.width = Pt(1.5)

        tb_f = s4.shapes.add_textbox(f_left + Inches(0.12), Inches(1.92), Inches(1.56), Inches(1.45))
        tf_f = tb_f.text_frame
        tf_f.word_wrap = True
        tf_f.margin_left = tf_f.margin_right = tf_f.margin_top = tf_f.margin_bottom = 0

        p_ft = tf_f.paragraphs[0]
        p_ft.text = f_title
        p_ft.font.size = Pt(10)
        p_ft.font.bold = True
        p_ft.font.color.rgb = f_col
        p_ft.space_after = Pt(6)

        p_fs = tf_f.add_paragraph()
        p_fs.text = f_sub
        p_fs.font.size = Pt(8.5)
        p_fs.font.color.rgb = COLOR_SLATE_TEXT
        p_fs.line_spacing = 1.15

        if f_idx < len(flow_steps) - 1:
            arr = s4.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, f_left + Inches(1.82), Inches(2.55), Inches(0.14), Inches(0.2))
            arr.fill.solid()
            arr.fill.fore_color.rgb = COLOR_SLATE_BORDER
            arr.line.color.rgb = COLOR_SLATE_BORDER

    # Bottom Half: The ErrorLog Database Table
    t_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.75), Inches(11.733), Inches(2.95))
    t_card.fill.solid()
    t_card.fill.fore_color.rgb = COLOR_WHITE
    t_card.line.color.rgb = COLOR_SLATE_BORDER
    t_card.line.width = Pt(1.5)

    tb_tlabel = s4.shapes.add_textbox(Inches(1.1), Inches(3.9), Inches(11.0), Inches(0.3))
    tf_tlabel = tb_tlabel.text_frame
    tf_tlabel.margin_left = tf_tlabel.margin_right = tf_tlabel.margin_top = tf_tlabel.margin_bottom = 0
    p_tl = tf_tlabel.paragraphs[0]
    p_tl.text = "DOCUMENTED ERRORLOG DATABASE SCHEMA (COLLECTION: errorlogs)"
    p_tl.font.size = Pt(10)
    p_tl.font.bold = True
    p_tl.font.color.rgb = COLOR_NAVY_PRIMARY

    # Add Table
    table_shape = s4.shapes.add_table(8, 4, Inches(1.1), Inches(4.25), Inches(11.133), Inches(2.25))
    table = table_shape.table
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(1.4)
    table.columns[2].width = Inches(2.2)
    table.columns[3].width = Inches(5.733)

    schema_rows = [
        ("Field Name", "Data Type", "Constraints / Default", "Description"),
        ("_id", "ObjectId", "Primary Key, Auto", "Unique telemetry incident identifier"),
        ("transactionId", "ObjectId", "Required, Ref: 'Transaction'", "Foreign reference to the affected ledger transaction"),
        ("errorCode", "String", "Required (e.g. INSUFFICIENT_FUNDS)", "Categorized machine-readable error designation"),
        ("severity", "String", "Enum: ['LOW','MEDIUM','HIGH','CRITICAL']", "Triage rating calculated from failure severity mapping"),
        ("message", "String", "Required, String", "Detailed human-readable failure context and diagnostics"),
        ("resolved", "Boolean", "Default: false", "Administrative resolution flag toggled upon investigation"),
        ("createdAt", "Date", "Auto Timestamp", "Exact ISO-8601 timestamp when failure was intercepted"),
    ]

    for r_i, row in enumerate(schema_rows):
        for c_i, val in enumerate(row):
            cell = table.cell(r_i, c_i)
            cell.text = val
            p_c = cell.text_frame.paragraphs[0]
            p_c.font.name = FONT_BODY
            if r_i == 0:
                p_c.font.size = Pt(9.5)
                p_c.font.bold = True
                p_c.font.color.rgb = COLOR_WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLOR_NAVY_PRIMARY
            else:
                p_c.font.size = Pt(8.5)
                p_c.font.color.rgb = COLOR_NAVY_DARK if c_i == 0 else COLOR_SLATE_TEXT
                if c_i == 0:
                    p_c.font.bold = True
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLOR_WHITE if r_i % 2 == 1 else COLOR_CARD_ALT

    # Speaker Notes Slide 4
    s4.notes_slide.notes_text_frame.text = (
        "Slide 4 details the Error Telemetry Pipeline in PayGate. Notice the sequential flowchart: "
        "When a transaction debit fails, a Mongoose post-save hook on the Transaction model "
        "automatically captures the failure reason, maps it to a severity tier, and creates an ErrorLog document in MongoDB. "
        "The schema shown below includes all 7 documented fields: transactionId, errorCode, severity, "
        "message, and the resolution boolean. Administrators inspect errors via GET /api/admin/errors and "
        "resolve them with PATCH /api/admin/errors/:id/resolve, guaranteeing end-to-end incident auditability."
    )

    # ════════════════════════════════════════════════════════════════════════════
    # SLIDE 5: Error Severity Classification
    # ════════════════════════════════════════════════════════════════════════════
    s5 = prs.slides.add_slide(blank_layout)
    bg5 = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg5.fill.solid()
    bg5.fill.fore_color.rgb = COLOR_BG_LIGHT
    bg5.line.color.rgb = COLOR_BG_LIGHT

    add_header_footer(s5, "Documented Error Severity Classification & Mapping", slide_num=5)

    severities = [
        ("CRITICAL", COLOR_SEV_CRIT,
         "Systemic or Partner Failures",
         [
             "Documented: BANK_TIMEOUT",
             "Documented: AUTH_FAILED",
             "Impact: Core banking partner unresponsive or security signature failure.",
             "Action: Immediate alert; requires gateway failover or partner escalation."
         ]),
        ("HIGH", COLOR_SEV_HIGH,
         "Financial & Security Breaches",
         [
             "Documented: INSUFFICIENT_FUNDS",
             "Documented: FROZEN_WALLET",
             "Documented: DUPLICATE_TXN / NETWORK_ERROR",
             "Impact: Customer attempts overdrawing balance or accessing frozen wallet.",
             "Action: Auto-blocks transfer; logged for audit review."
         ]),
        ("MEDIUM", COLOR_SEV_MED,
         "Validation & Limit Breaches",
         [
             "Documented: INVALID_ACCOUNT",
             "Documented: LIMIT_EXCEEDED",
             "Impact: Destination bank IFSC/account mismatch or daily cap exceeded.",
             "Action: Client alerted to update beneficiary details."
         ]),
        ("LOW", COLOR_SEV_LOW,
         "Informational Warnings",
         [
             "Documented: Informational / Minor Warnings",
             "Impact: Non-critical anomalies or routine operational notifications.",
             "Action: Logged for historical telemetry without requiring active triage."
         ]),
    ]

    for s_idx, (s_name, s_col, s_desc, s_bullets) in enumerate(severities):
        s_left = Inches(0.8 + s_idx * 3.0)
        scard = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, s_left, Inches(1.8), Inches(2.78), Inches(4.9))
        scard.fill.solid()
        scard.fill.fore_color.rgb = COLOR_WHITE
        scard.line.color.rgb = COLOR_SLATE_BORDER
        scard.line.width = Pt(1.5)

        # Header banner inside card
        s_banner = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, s_left, Inches(1.8), Inches(2.78), Inches(0.75))
        s_banner.fill.solid()
        s_banner.fill.fore_color.rgb = s_col
        s_banner.line.color.rgb = s_col

        tb_sh = s5.shapes.add_textbox(s_left, Inches(1.95), Inches(2.78), Inches(0.5))
        tf_sh = tb_sh.text_frame
        tf_sh.margin_left = tf_sh.margin_right = tf_sh.margin_top = tf_sh.margin_bottom = 0
        p_sh = tf_sh.paragraphs[0]
        p_sh.alignment = PP_ALIGN.CENTER
        p_sh.text = s_name
        p_sh.font.size = Pt(14)
        p_sh.font.bold = True
        p_sh.font.color.rgb = COLOR_WHITE

        # Body text
        tb_sb = s5.shapes.add_textbox(s_left + Inches(0.2), Inches(2.7), Inches(2.38), Inches(3.8))
        tf_sb = tb_sb.text_frame
        tf_sb.word_wrap = True
        tf_sb.margin_left = tf_sb.margin_right = tf_sb.margin_top = tf_sb.margin_bottom = 0

        p_sdesc = tf_sb.paragraphs[0]
        p_sdesc.text = s_desc.upper()
        p_sdesc.font.size = Pt(9)
        p_sdesc.font.bold = True
        p_sdesc.font.color.rgb = s_col
        p_sdesc.space_after = Pt(12)

        for sb in s_bullets:
            p_sb = tf_sb.add_paragraph()
            p_sb.text = f"• {sb}"
            p_sb.font.size = Pt(9.5)
            p_sb.font.color.rgb = COLOR_SLATE_TEXT
            p_sb.space_after = Pt(8)
            p_sb.line_spacing = 1.15

    # Speaker Notes Slide 5
    s5.notes_slide.notes_text_frame.text = (
        "Slide 5 outlines the four documented error severity tiers in PayGate: CRITICAL, HIGH, MEDIUM, and LOW. "
        "In our codebase, the Mongoose post-save hook explicitly maps failure reasons: "
        "BANK_TIMEOUT and AUTH_FAILED are flagged as CRITICAL because partner connectivity is affected. "
        "INSUFFICIENT_FUNDS, FROZEN_WALLET, and DUPLICATE_TXN are categorized as HIGH because they involve "
        "account balance integrity and fraud prevention. "
        "MEDIUM handles account and limit mismatches, while LOW records informational warnings."
    )

    # ════════════════════════════════════════════════════════════════════════════
    # SLIDE 6: Pending Transactions & Operational Recovery
    # ════════════════════════════════════════════════════════════════════════════
    s6 = prs.slides.add_slide(blank_layout)
    bg6 = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg6.fill.solid()
    bg6.fill.fore_color.rgb = COLOR_BG_LIGHT
    bg6.line.color.rgb = COLOR_BG_LIGHT

    add_header_footer(s6, "Pending Transaction Monitoring & Administrative Recovery", slide_num=6)

    # Left Column: Recovery Role & Lifecycle
    p_left_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9))
    p_left_card.fill.solid()
    p_left_card.fill.fore_color.rgb = COLOR_WHITE
    p_left_card.line.color.rgb = COLOR_SLATE_BORDER
    p_left_card.line.width = Pt(1.5)

    tb_pl = s6.shapes.add_textbox(Inches(1.1), Inches(2.05), Inches(5.0), Inches(4.4))
    tf_pl = tb_pl.text_frame
    tf_pl.word_wrap = True
    tf_pl.margin_left = tf_pl.margin_right = tf_pl.margin_top = tf_pl.margin_bottom = 0

    p_plt = tf_pl.paragraphs[0]
    p_plt.text = "ADMINISTRATIVE RECOVERY ROLE"
    p_plt.font.size = Pt(10)
    p_plt.font.bold = True
    p_plt.font.color.rgb = COLOR_BLUE_ACCENT

    p_plh = tf_pl.add_paragraph()
    p_plh.text = "In-Flight Clearing & Recovery Controls"
    p_plh.font.size = Pt(14)
    p_plh.font.bold = True
    p_plh.font.color.rgb = COLOR_NAVY_DARK
    p_plh.space_after = Pt(14)

    role_points = [
        "In-Flight Transaction Queue: Payments in status: 'pending' indicate active inter-bank clearance or pending OTP verification.",
        "Gateway Time-to-Live (TTL): Tracks elapsed time against a maximum 30-minute threshold. Flags transactions nearing timeout in amber/red.",
        "Manual Settlement Simulation: Allows authorized operators to confirm settlement if the external banking acknowledgment was delayed.",
        "Defensive Cancellation: Allows operators to cancel stalled transactions, safely unlocking reserved customer funds."
    ]
    for rp in role_points:
        p_rp = tf_pl.add_paragraph()
        p_rp.text = f"• {rp}"
        p_rp.font.size = Pt(10)
        p_rp.font.color.rgb = COLOR_SLATE_TEXT
        p_rp.space_after = Pt(10)
        p_rp.line_spacing = 1.15

    # Right Column: Workflow / Architecture
    p_right_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.9))
    p_right_card.fill.solid()
    p_right_card.fill.fore_color.rgb = COLOR_WHITE
    p_right_card.line.color.rgb = COLOR_SLATE_BORDER
    p_right_card.line.width = Pt(1.5)

    tb_pr = s6.shapes.add_textbox(Inches(7.1), Inches(2.05), Inches(5.133), Inches(4.4))
    tf_pr = tb_pr.text_frame
    tf_pr.word_wrap = True
    tf_pr.margin_left = tf_pr.margin_right = tf_pr.margin_top = tf_pr.margin_bottom = 0

    p_prt = tf_pr.paragraphs[0]
    p_prt.text = "OPERATIONAL RESOLUTION WORKFLOW"
    p_prt.font.size = Pt(10)
    p_prt.font.bold = True
    p_prt.font.color.rgb = COLOR_TEAL_ACCENT

    p_prh = tf_pr.add_paragraph()
    p_prh.text = "Exact API Endpoints in Codebase"
    p_prh.font.size = Pt(14)
    p_prh.font.bold = True
    p_prh.font.color.rgb = COLOR_NAVY_DARK
    p_prh.space_after = Pt(14)

    # 3 Sub-blocks for actions with exact project routes
    action_blocks = [
        ("Action A: Force Settlement (Success)", 
         "Endpoint: PATCH /api/admin/transactions/:id/force-success\n"
         "Action: Updates status to 'success', records admin in resolvedBy, and commits final ledger settlement.", 
         COLOR_TEAL_ACCENT),
        ("Action B: Force Failure / Cancel (Terminal State)", 
         "Endpoint: PATCH /api/admin/transactions/:id/force-fail\n"
         "Action: Updates status to 'failure' with reason 'BANK_TIMEOUT', creates ErrorLog, and halts in-flight hang.", 
         COLOR_SEV_CRIT),
        ("Action C: Wallet Security Freeze & Account Control", 
         "Endpoint: PATCH /api/admin/users/:id/freeze\n"
         "Action: Toggles wallet.isFrozen instantly, blocking outgoing debits if suspicious activity is identified.", 
         COLOR_NAVY_PRIMARY)
    ]
    for act_t, act_d, act_c in action_blocks:
        p_at = tf_pr.add_paragraph()
        p_at.text = act_t
        p_at.font.size = Pt(11)
        p_at.font.bold = True
        p_at.font.color.rgb = act_c

        p_ad = tf_pr.add_paragraph()
        p_ad.text = act_d
        p_ad.font.size = Pt(9.5)
        p_ad.font.color.rgb = COLOR_SLATE_TEXT
        p_ad.space_after = Pt(10)
        p_ad.line_spacing = 1.15

    # Speaker Notes Slide 6
    s6.notes_slide.notes_text_frame.text = (
        "Slide 6 covers Pending Transaction Resolution and Operational Recovery in PayGate. "
        "Transactions in 'pending' status represent in-flight payments awaiting clearing from partner banks. "
        "The Admin Console features a 30-minute TTL timeout monitor. "
        "If bank clearing is delayed, the administrator can perform two documented actions in our codebase: "
        "First, 'Force Settlement' via PATCH /api/admin/transactions/:id/force-success to transition the payment to 'success' once cleared out-of-band. "
        "Or second, 'Force Failure' via PATCH /api/admin/transactions/:id/force-fail to prevent indefinite hangs and release customer funds. "
        "Additionally, the administrator can trigger an immediate wallet security freeze via PATCH /api/admin/users/:id/freeze if fraud is suspected."
    )

    # ════════════════════════════════════════════════════════════════════════════
    # SLIDE 7: Conclusion & Value Matrix
    # ════════════════════════════════════════════════════════════════════════════
    s7 = prs.slides.add_slide(blank_layout)
    bg7 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg7.fill.solid()
    bg7.fill.fore_color.rgb = COLOR_BG_LIGHT
    bg7.line.color.rgb = COLOR_BG_LIGHT

    add_header_footer(s7, "Conclusion: Architectural Value of the Admin Console", slide_num=7)

    # 5 Pillars of Enterprise Value
    pillars_5 = [
        ("1. Real-Time Visibility", 
         "Immediate operational telemetry into system-wide volume, success percentages, and database heartbeat every 10 seconds.",
         COLOR_NAVY_PRIMARY),
        ("2. Complete Auditability",
         "Every financial debit and administrative resolution is permanently anchored in MongoDB with cryptographic reference codes.",
         COLOR_BLUE_ACCENT),
        ("3. Error Traceability",
         "Automated capture of failure reasons into the ErrorLog collection with 4-tier severity triage (LOW to CRITICAL).",
         COLOR_SEV_HIGH),
        ("4. Operational Control",
         "Zero-delay wallet freeze kill-switch (PATCH /api/admin/users/:id/freeze) preventing unauthorized balance drains during active security incidents.",
         COLOR_SEV_CRIT),
        ("5. Rapid Recovery Support",
         "Deterministic pending transaction resolution (force-success / force-fail) preventing stuck payments and eliminating reconciliation overhead.",
         COLOR_TEAL_ACCENT),
    ]

    for p_idx, (p_t, p_d, p_c) in enumerate(pillars_5):
        top_y = Inches(1.75 + p_idx * 1.0)
        card_p = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_y, Inches(11.733), Inches(0.88))
        card_p.fill.solid()
        card_p.fill.fore_color.rgb = COLOR_WHITE
        card_p.line.color.rgb = COLOR_SLATE_BORDER
        card_p.line.width = Pt(1.5)

        side_p = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top_y, Inches(0.12), Inches(0.88))
        side_p.fill.solid()
        side_p.fill.fore_color.rgb = p_c
        side_p.line.color.rgb = p_c

        tb_p = s7.shapes.add_textbox(Inches(1.1), top_y + Inches(0.1), Inches(11.2), Inches(0.68))
        tf_p = tb_p.text_frame
        tf_p.word_wrap = True
        tf_p.margin_left = tf_p.margin_right = tf_p.margin_top = tf_p.margin_bottom = 0

        p_pt = tf_p.paragraphs[0]
        p_pt.text = p_t
        p_pt.font.size = Pt(11.5)
        p_pt.font.bold = True
        p_pt.font.color.rgb = COLOR_NAVY_DARK

        p_pd = tf_p.add_paragraph()
        p_pd.text = p_d
        p_pd.font.size = Pt(9.5)
        p_pd.font.color.rgb = COLOR_SLATE_TEXT
        p_pd.space_before = Pt(2)

    # Speaker Notes Slide 7
    s7.notes_slide.notes_text_frame.text = (
        "In conclusion, Module 6 fulfills the crucial operational and audit requirements of PayGate. "
        "It provides real-time visibility, rigorous auditability, automated error traceability, "
        "decisive operational control via wallet freeze triggers, and active recovery support for pending transfers. "
        "This ensures that PayGate operates not just as a consumer payment wallet, but as an enterprise-grade, "
        "fault-tolerant financial settlement gateway. "
        "Thank you, and I now welcome any questions regarding Module 6."
    )

    out_path = r"D:\Clg\python lab\MERN PROJECT\PayGate_Module6_Admin_Console_Presentation.pptx"
    updated_path = r"D:\Clg\python lab\MERN PROJECT\PayGate_Module6_Admin_Console_Presentation_v2.pptx"
    try:
        prs.save(out_path)
        print(f"SUCCESS: Presentation created at {out_path}")
    except PermissionError:
        prs.save(updated_path)
        print(f"File currently open in PowerPoint. SUCCESS: Updated presentation created at {updated_path}")

if __name__ == '__main__':
    create_presentation()
