import docx
import sys
import xml.etree.ElementTree as ET

sys.stdout.reconfigure(encoding='utf-8')

ref_path = r'C:\Users\dell\Downloads\ref.docx'
doc = docx.Document(ref_path)

print(f"=== DOCUMENT METRICS ===")
print(f"Paragraphs: {len(doc.paragraphs)}")
print(f"Tables: {len(doc.tables)}")
print(f"Sections: {len(doc.sections)}")

sec = doc.sections[0]
print(f"Page Width: {sec.page_width.inches} in ({sec.page_width.pt} pt)")
print(f"Page Height: {sec.page_height.inches} in ({sec.page_height.pt} pt)")
print(f"Top Margin: {sec.top_margin.inches} in ({sec.top_margin.pt} pt)")
print(f"Bottom Margin: {sec.bottom_margin.inches} in ({sec.bottom_margin.pt} pt)")
print(f"Left Margin: {sec.left_margin.inches} in ({sec.left_margin.pt} pt)")
print(f"Right Margin: {sec.right_margin.inches} in ({sec.right_margin.pt} pt)")

print("\n=== PARAGRAPH BY PARAGRAPH BREAKDOWN ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    runs_details = []
    for r in p.runs:
        c = str(r.font.color.rgb) if r.font.color and r.font.color.rgb else 'None'
        sz = r.font.size.pt if r.font.size else 'None'
        fn = r.font.name
        b = r.font.bold
        it = r.font.italic
        runs_details.append(f"[{fn}, {sz}pt, b={b}, i={it}, clr={c}: '{r.text}']")
    
    pf = p.paragraph_format
    sb = pf.space_before.pt if pf.space_before else 'None'
    sa = pf.space_after.pt if pf.space_after else 'None'
    ls = pf.line_spacing
    al = p.alignment
    has_img = '<w:drawing' in ET.tostring(p._p, encoding='unicode')
    
    t_short = t.encode('ascii', 'replace').decode('ascii')[:50]
    print(f"P[{i:3d}] al={str(al):10s} sb={str(sb):5s} sa={str(sa):5s} ls={str(ls):5s} img={str(has_img):5s} | '{t_short}'")
    if runs_details:
        print(f"       runs: {' '.join(runs_details[:4])}")
