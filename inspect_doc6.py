import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc_path = r'C:\Users\dell\Desktop\PayGate_Module_6_Professionally_Formatted.docx'
doc = docx.Document(doc_path)

style_norm = doc.styles['Normal']
print(f"Normal style font: name={style_norm.font.name}, size={style_norm.font.size.pt if style_norm.font.size else None}")

for i in range(13):
    p = doc.paragraphs[i]
    r_info = []
    for r in p.runs:
        t_clean = r.text.encode('ascii', 'replace').decode('ascii')[:20]
        r_info.append((t_clean, r.font.name, r.font.size.pt if r.font.size else None, r.font.bold))
    print(f"P[{i:2d}]: {r_info}")


print(f"Total paragraphs: {len(doc.paragraphs)}")

for i in range(14, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    t = p.text.strip()
    xml_str = p._element.xml
    has_pb = 'w:br w:type="page"' in xml_str or 'w:pageBreakBefore' in xml_str
    
    # check if heading or bold
    is_heading = False
    for r in p.runs:
        if r.font.bold and r.font.size and r.font.size.pt >= 10.5:
            is_heading = True
            break
            
    # Also check short text or section titles
    if is_heading or has_pb or (len(t) > 0 and len(t) < 70 and not t.startswith('•') and not t.startswith('-')):
        t_clean = t.encode('ascii', 'replace').decode('ascii')[:60]
        print(f"P[{i:3d}] pb={has_pb} text='{t_clean}'")




