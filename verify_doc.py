import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'C:\Users\dell\Desktop\PayGate_Module_6_Professionally_Formatted.docx')

print("=== VERIFYING COVER PAGE ===")
p0 = doc.paragraphs[0]
print(f"P0 (Logo): alignment={p0.alignment}, space_after={p0.paragraph_format.space_after.pt if p0.paragraph_format.space_after else None}")
p1 = doc.paragraphs[1]
print(f"P1 (College Name): alignment={p1.alignment}, text='{p1.text.strip()}'")

print("\n=== VERIFYING TABLE HEADERS (WHITE FONT) ===")
for idx, tbl in enumerate(doc.tables):
    row0 = tbl.rows[0]
    header_texts = [c.text.strip() for c in row0.cells]
    header_colors = []
    for c in row0.cells:
        cell_colors = []
        for p in c.paragraphs:
            for r in p.runs:
                cell_colors.append((r.text.strip(), str(r.font.color.rgb) if r.font.color else 'None'))
        header_colors.append(cell_colors)
    print(f"Table {idx}:")
    print(f"  Headers: {header_texts}")
    print(f"  Colors:  {header_colors}")

print("\n=== VERIFYING MAJOR HEADINGS & PAGE BREAKS ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    xml = p._element.xml
    has_page_break_before = 'w:pageBreakBefore' in xml
    has_br = 'w:br w:type="page"' in xml
    keep = p.paragraph_format.keep_with_next
    if any(t == s or t.startswith(s) for s in [
        "Synopsis", "Objective", "Scope", "Project Description", "DB Design", 
        "5.2 transactions", "5.3 Documented", "Data Flow Diagram", "ER Diagram",
        "Use Case Diagram", "Class Diagram", "Sequence Diagram", "Screenshots (Forms)",
        "Form 12.1", "Form 12.3", "Form 12.5", "Test Report", "Conclusion and Further"
    ]):
        print(f"P[{i:3d}] break_before={has_page_break_before} br={has_br} keep={keep} | '{t[:50]}'")
