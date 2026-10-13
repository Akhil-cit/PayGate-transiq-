import docx, os, sys
sys.stdout.reconfigure(encoding='utf-8')

all_dir = r'C:\Users\dell\Downloads\all 6'
for fname in ['1.docx', '2.docx', '3.docx', '4.docx', '5.docx']:
    doc = docx.Document(os.path.join(all_dir, fname))
    print(f'=== Image placements in {fname} ===')
    for p_idx, p in enumerate(doc.paragraphs):
        for r in p.runs:
            if 'graphic' in r._element.xml:
                prev_t = doc.paragraphs[p_idx-1].text.strip() if p_idx > 0 else ''
                curr_t = p.text.strip()
                next_t = doc.paragraphs[p_idx+1].text.strip() if p_idx+1 < len(doc.paragraphs) else ''
                print(f'  P[{p_idx}]: prev="{prev_t[:45]}" | next="{next_t[:45]}"')
