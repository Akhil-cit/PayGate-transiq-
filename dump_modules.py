import docx, os, sys
sys.stdout.reconfigure(encoding='utf-8')

all_dir = r'C:\Users\dell\Downloads\all 6'
os.makedirs('module_dumps', exist_ok=True)

for fname in ['1.docx', '2.docx', '3.docx', '4.docx', '5.docx']:
    fpath = os.path.join(all_dir, fname)
    doc = docx.Document(fpath)
    out_name = os.path.join('module_dumps', f'{fname}.txt')
    with open(out_name, 'w', encoding='utf-8') as f:
        f.write(f'=== {fname} ===\n\n')
        for i, p in enumerate(doc.paragraphs):
            t = p.text.strip()
            if t:
                f.write(f'[P{i}] ({p.style.name}) {t}\n')
        f.write('\n\n=== TABLES ===\n')
        for t_idx, t in enumerate(doc.tables):
            f.write(f'\n--- Table {t_idx} ({len(t.rows)}x{len(t.columns)}) ---\n')
            for r_idx, r in enumerate(t.rows):
                row_t = [c.text.strip().replace('\n', ' ') for c in r.cells]
                f.write(f'Row {r_idx}: ' + ' | '.join(row_t) + '\n')
print('Dumped 1.docx to 5.docx successfully!')
