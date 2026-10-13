import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os
import shutil

# Color palette
NAVY = RGBColor(15, 39, 71)      # 0F2747
SLATE = RGBColor(28, 70, 112)    # 1C4670
TEAL = RGBColor(13, 148, 136)    # 0D9488
DARK = RGBColor(38, 48, 60)      # 26303C
MUTED = RGBColor(100, 116, 139)  # 64748B
WHITE = RGBColor(255, 255, 255)

HEX_NAVY = "0F2747"
HEX_SLATE = "1C4670"
HEX_LIGHT_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"

def set_cell_margins(cell, top=65, bottom=65, left=110, right=110):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin, value in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin}')
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_table_borders(table, color="CBD5E1", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def style_table(table, col_widths, col_alignments=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for r_idx, row in enumerate(table.rows):
        is_header = (r_idx == 0)
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for c_idx, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[c_idx])
            set_cell_margins(cell, top=65, bottom=65, left=110, right=110)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            if is_header:
                set_cell_shading(cell, HEX_NAVY)
            else:
                if r_idx % 2 == 1:
                    set_cell_shading(cell, "FFFFFF")
                else:
                    set_cell_shading(cell, HEX_LIGHT_BG)
                    
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                if col_alignments and c_idx < len(col_alignments):
                    p.alignment = col_alignments[c_idx]
                for r in p.runs:
                    r.font.name = 'Aptos'
                    if is_header:
                        r.font.bold = True
                        r.font.size = Pt(9)
                        r.font.color.rgb = WHITE
                    else:
                        r.font.size = Pt(8.5)
                        r.font.color.rgb = DARK

class DocBuilder:
    def __init__(self):
        self.doc = docx.Document()
        for s in self.doc.sections:
            s.page_width = Inches(8.5)
            s.page_height = Inches(11.0)
            s.top_margin = Inches(0.65)
            s.bottom_margin = Inches(0.65)
            s.left_margin = Inches(0.75)
            s.right_margin = Inches(0.75)
            
        style_normal = self.doc.styles['Normal']
        style_normal.font.name = 'Aptos'
        style_normal.font.size = Pt(10)
        style_normal.font.color.rgb = DARK
        style_normal.paragraph_format.line_spacing = 1.12
        style_normal.paragraph_format.space_after = Pt(3)

    def add_h1(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Aptos Display'
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = NAVY
        return p

    def add_h2(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Aptos'
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = SLATE
        return p

    def add_h3(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Aptos'
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = TEAL
        return p

    def add_p(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        r.font.name = 'Aptos'
        r.font.size = Pt(10)
        r.font.color.rgb = DARK
        return p

    def add_bullet(self, text, bold_prefix=""):
        p = self.doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.12
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Aptos'
            r_pre.font.size = Pt(9.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = NAVY
        r = p.add_run(text)
        r.font.name = 'Aptos'
        r.font.size = Pt(9.5)
        r.font.color.rgb = DARK
        return p

    def add_caption(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.font.name = 'Aptos'
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.bold = True
        r.font.color.rgb = SLATE
        return p

    def add_image(self, path, width_in=6.0, caption=""):
        if os.path.exists(path):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.add_run().add_picture(path, width=Inches(width_in))
            if caption:
                self.add_caption(caption)
        else:
            print(f"Warning: Image not found {path}")

    def add_table_data(self, headers, rows_data, col_widths, col_alignments=None):
        table = self.doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        for c_i, h in enumerate(headers):
            table.rows[0].cells[c_i].paragraphs[0].text = h
        for r_i, r_data in enumerate(rows_data):
            for c_i, val in enumerate(r_data):
                table.rows[r_i + 1].cells[c_i].paragraphs[0].text = str(val)
        style_table(table, col_widths, col_alignments)
        return table

    def build_cover_page(self, mod_num, mod_title, student_name, student_reg):
        # Logo
        p_logo = self.doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(6)
        p_logo.paragraph_format.space_after = Pt(6)
        if os.path.exists('cit_logo.png'):
            p_logo.add_run().add_picture('cit_logo.png', width=Inches(1.15))

        # College Name
        p_clg = self.doc.add_paragraph()
        p_clg.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_clg.paragraph_format.space_after = Pt(1)
        r_clg = p_clg.add_run("COIMBATORE INSTITUTE OF TECHNOLOGY")
        r_clg.font.name = 'Aptos Display'
        r_clg.font.size = Pt(14.5)
        r_clg.font.bold = True
        r_clg.font.color.rgb = NAVY

        p_sub = self.doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(8)
        r_sub = p_sub.add_run("(Government Aided Autonomous Institution affiliated to Anna University)\nCoimbatore – 641 014, Tamil Nadu, India")
        r_sub.font.name = 'Aptos'
        r_sub.font.size = Pt(9)
        r_sub.font.color.rgb = SLATE

        p_bar = self.doc.add_paragraph()
        p_bar.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_bar.paragraph_format.space_after = Pt(8)
        r_bar = p_bar.add_run("━" * 56)
        r_bar.font.color.rgb = NAVY
        r_bar.font.size = Pt(9.5)

        p_rec = self.doc.add_paragraph()
        p_rec.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_rec.paragraph_format.space_after = Pt(2)
        r_rec = p_rec.add_run("PROJECT DOCUMENTATION RECORD")
        r_rec.font.name = 'Aptos Display'
        r_rec.font.size = Pt(11)
        r_rec.font.bold = True
        r_rec.font.color.rgb = TEAL

        p_title = self.doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(3)
        r_title = p_title.add_run("PayGate (Transiq)")
        r_title.font.name = 'Aptos Display'
        r_title.font.size = Pt(20)
        r_title.font.bold = True
        r_title.font.color.rgb = NAVY

        p_subtitle = self.doc.add_paragraph()
        p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_subtitle.paragraph_format.space_after = Pt(3)
        r_subtitle = p_subtitle.add_run("Secure Digital Payment & Wallet Management Platform")
        r_subtitle.font.name = 'Aptos'
        r_subtitle.font.size = Pt(12)
        r_subtitle.font.bold = True
        r_subtitle.font.color.rgb = SLATE

        p_sub2 = self.doc.add_paragraph()
        p_sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub2.paragraph_format.space_after = Pt(8)
        r_sub2 = p_sub2.add_run("A Scalable Full-Stack MERN Payment Gateway and Intelligent Digital Wallet System")
        r_sub2.font.name = 'Aptos'
        r_sub2.font.size = Pt(10)
        r_sub2.font.italic = True
        r_sub2.font.color.rgb = MUTED

        p_mod = self.doc.add_paragraph()
        p_mod.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_mod.paragraph_format.space_after = Pt(10)
        r_mod = p_mod.add_run(f"MODULE {mod_num}: {mod_title.upper()}")
        r_mod.font.name = 'Aptos Display'
        r_mod.font.size = Pt(12)
        r_mod.font.bold = True
        r_mod.font.color.rgb = NAVY

        # Project Details
        p_dth = self.doc.add_paragraph()
        p_dth.paragraph_format.space_after = Pt(3)
        r_dth = p_dth.add_run("PROJECT DETAILS")
        r_dth.font.name = 'Aptos'
        r_dth.font.size = Pt(10.5)
        r_dth.font.bold = True
        r_dth.font.color.rgb = NAVY

        t_data = [
            ("Project Title", "PayGate (Transiq) — Secure Digital Payment & Wallet Management Platform"),
            ("Subsystem Focus", f"Module {mod_num}: {mod_title}"),
            ("Domain", "Full-Stack Web Development — Enterprise FinTech & Digital Payments (MERN Stack)"),
            ("Institution", "Coimbatore Institute of Technology (Autonomous), Coimbatore – 641 014"),
            ("Course & Degree", "M.Sc. Software Systems (5-Year Integrated Degree Programme)"),
            ("Academic Year", "2025 – 2026 (Semester III)")
        ]
        self.add_table_data(["Field", "Details"], t_data, [2.2, 4.8], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

        # Team Details
        p_tmh = self.doc.add_paragraph()
        p_tmh.paragraph_format.space_before = Pt(8)
        p_tmh.paragraph_format.space_after = Pt(3)
        r_tmh = p_tmh.add_run("TEAM DETAILS (6 MEMBERS)")
        r_tmh.font.name = 'Aptos'
        r_tmh.font.size = Pt(10.5)
        r_tmh.font.bold = True
        r_tmh.font.color.rgb = NAVY

        team_rows = [
            ("1", "Krishnaprabha D", "2503717672622025", "Module 1: Customer Profile, Wallets & Smart Pockets" + (" (Lead)" if mod_num==1 else "")),
            ("2", "Nithinkumar R", "2503717672621035", "Module 2: Beneficiary Management & Directory Services" + (" (Lead)" if mod_num==2 else "")),
            ("3", "Mohamed Riyaz A", "2503717672621028", "Module 3: Idempotent Payment Initiation & Checkout Wizard" + (" (Lead)" if mod_num==3 else "")),
            ("4", "Rithick K", "2503717672621040", "Module 4: Transaction Ledgers & Financial Analytics" + (" (Lead)" if mod_num==4 else "")),
            ("5", "Mohamed Umar K", "2503717672621029", "Module 5: Security Governance, Access Control & Wallet Kill-Switch" + (" (Lead)" if mod_num==5 else "")),
            ("6", "Dinesh G", "2503717672621008", "Module 6: Admin Operations, Monitoring & Error Audit Console" + (" (Lead)" if mod_num==6 else ""))
        ]
        self.add_table_data(["S.No", "Name of Student", "Register Number", "Assigned Module / Role"], team_rows, [0.6, 2.2, 1.6, 2.6], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])

        # Signatures
        p_gd = self.doc.add_paragraph()
        p_gd.paragraph_format.space_before = Pt(10)
        p_gd.paragraph_format.space_after = Pt(3)
        p_gd.add_run("Project Guide / Supervisor: Prof. / Dr. __________________________________________________").font.color.rgb = DARK

        p_sig = self.doc.add_paragraph()
        p_sig.paragraph_format.space_before = Pt(4)
        p_sig.paragraph_format.space_after = Pt(0)
        p_sig.add_run("Signature of the Guide: ______________________________      Date of Submission: _______________").font.color.rgb = DARK

        self.doc.add_page_break()

    def build_conclusion_enhancements(self, conclusion_text, enhancements_list):
        self.add_h1("Conclusion and Further Enhancements")
        self.add_h2("Conclusion")
        for p_t in conclusion_text:
            self.add_p(p_t)
        self.add_h2("Further Enhancements")
        for pre, t in enhancements_list:
            self.add_bullet(t, pre)
            
        p_end = self.doc.add_paragraph()
        p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_end.paragraph_format.space_before = Pt(14)
        r_end = p_end.add_run("— END OF DOCUMENTATION RECORD —")
        r_end.font.name = 'Aptos'
        r_end.font.size = Pt(9.5)
        r_end.font.bold = True
        r_end.font.color.rgb = MUTED

    def save(self, target_path):
        self.doc.save(target_path)
        print(f"Saved: {target_path}")

print("DocBuilder class loaded successfully.")
