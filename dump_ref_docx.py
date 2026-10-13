import docx
import sys
import xml.etree.ElementTree as ET

sys.stdout.reconfigure(encoding='utf-8')

ref_path = r'C:\Users\dell\Downloads\ref.docx'
doc = docx.Document(ref_path)

print(f"Total sections: {len(doc.sections)}")
for i, sec in enumerate(doc.sections):
    print(f"Section {i}: page_w={sec.page_width.inches} in, page_h={sec.page_height.inches} in")
    print(f"  top={sec.top_margin.inches} in, bottom={sec.bottom_margin.inches} in, left={sec.left_margin.inches} in, right={sec.right_margin.inches} in")
    print(f"  header_dist={sec.header_distance.inches} in, footer_dist={sec.footer_distance.inches} in")
    print(f"  different_first_page={sec.different_first_page_header_footer}")
    sectPr = sec._sectPr
    for el in sectPr:
        tag = el.tag.split('}')[-1]
        print(f"    sectPr child: {tag} -> {ET.tostring(el, encoding='unicode')[:150]}")

print("\n--- ALL PARAGRAPHS IN REF.DOCX ---")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    xml_str = ET.tostring(p._p, encoding='unicode')
    has_pb = ('w:br w:type="page"' in xml_str) or ('w:pageBreakBefore' in xml_str)
    has_img = '<w:drawing' in xml_str
    
    # Check borders on paragraph
    pBdr = p._p.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr/{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr')
    bdr_info = ET.tostring(pBdr, encoding='unicode')[:50] if pBdr is not None else 'None'
    
    runs_info = []
    for r in p.runs:
        c = str(r.font.color.rgb) if r.font.color and r.font.color.rgb else 'None'
        runs_info.append((r.text[:15], r.font.name, r.font.size.pt if r.font.size else None, r.font.bold, c))
    
    al = p.alignment
    sb = p.paragraph_format.space_before.pt if p.paragraph_format.space_before else None
    sa = p.paragraph_format.space_after.pt if p.paragraph_format.space_after else None
    ls = p.paragraph_format.line_spacing
    
    t_clean = t.encode('ascii', 'replace').decode('ascii')[:40]
    print(f"P[{i:3d}] style='{p.style.name:12s}' al={str(al):10s} sb={str(sb):4s} sa={str(sa):4s} ls={str(ls):4s} pb={has_pb} img={has_img} bdr={bdr_info[:15]} | '{t_clean}' | runs={runs_info[:3]}")

print("\n--- ALL TABLES IN REF.DOCX ---")
for t_idx, tbl in enumerate(doc.tables):
    tblBdr = tbl._tbl.tblPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblBorders')
    bdr_str = ET.tostring(tblBdr, encoding='unicode') if tblBdr is not None else 'None'
    print(f"\nTable {t_idx}: rows={len(tbl.rows)}, cols={len(tbl.columns)}, style={tbl.style.name}, borders={bdr_str[:120]}")
    for r_idx, row in enumerate(tbl.rows):
        cells_info = []
        for c_idx, cell in enumerate(row.cells):
            shd = cell._tc.get_or_add_tcPr().find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd')
            fill = shd.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill') if shd is not None else None
            # get font info
            c_runs = []
            for cp in cell.paragraphs:
                for cr in cp.runs:
                    cc = str(cr.font.color.rgb) if cr.font.color and cr.font.color.rgb else 'None'
                    c_runs.append((cr.text[:10], cr.font.name, cr.font.size.pt if cr.font.size else None, cr.font.bold, cc))
            cells_info.append((cell.text[:15].strip(), fill, c_runs[:2]))
        print(f"  Row {r_idx}: {cells_info[:2]}")
