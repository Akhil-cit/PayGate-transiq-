import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'C:\Users\dell\Downloads\ref.docx')

print("=== PARAGRAPHS 0 to 18 IN REF.DOCX ===")
for i in range(19):
    p = doc.paragraphs[i]
    runs = [(r.text[:30], r.font.name, r.font.size.pt if r.font.size else None, r.font.bold, r.font.italic, str(r.font.color.rgb) if r.font.color and r.font.color.rgb else 'None') for r in p.runs]
    sb = p.paragraph_format.space_before.pt if p.paragraph_format.space_before else None
    sa = p.paragraph_format.space_after.pt if p.paragraph_format.space_after else None
    ls = p.paragraph_format.line_spacing
    al = p.alignment
    has_img = '<w:drawing' in p._element.xml
    print(f"P[{i:2d}] al={str(al):10s} sb={str(sb):4s} sa={str(sa):4s} ls={str(ls):4s} img={has_img} | text='{p.text[:40]}'")
    print(f"     runs={runs}")

print("\n=== MARGINS IN REF.DOCX ===")
sec = doc.sections[0]
print(f"Top: {sec.top_margin.inches} in ({sec.top_margin.pt} pt)")
print(f"Bottom: {sec.bottom_margin.inches} in ({sec.bottom_margin.pt} pt)")
print(f"Left: {sec.left_margin.inches} in ({sec.left_margin.pt} pt)")
print(f"Right: {sec.right_margin.inches} in ({sec.right_margin.pt} pt)")
print(f"Page Width: {sec.page_width.inches} in")
print(f"Page Height: {sec.page_height.inches} in")

print("\n=== TABLE STYLES & ROW COLORS IN REF.DOCX ===")
for t_idx, tbl in enumerate(doc.tables):
    print(f"Table {t_idx}: rows={len(tbl.rows)}, cols={len(tbl.columns)}")
    for r_idx in range(min(3, len(tbl.rows))):
        row = tbl.rows[r_idx]
        row_data = []
        for cell in row.cells:
            shd = cell._tc.get_or_add_tcPr().find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd')
            fill = shd.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill') if shd is not None else None
            c_text = cell.text.strip()[:15]
            c_runs = [(r.text[:10], r.font.name, r.font.size.pt if r.font.size else None, r.font.bold, str(r.font.color.rgb) if r.font.color and r.font.color.rgb else 'None') for p in cell.paragraphs for r in p.runs]
            row_data.append((c_text, fill, c_runs))
        print(f"  Row {r_idx}: {row_data}")
