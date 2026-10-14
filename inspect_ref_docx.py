import docx
import sys
from docx.oxml import parse_xml
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')

ref_path = r'C:\Users\dell\Downloads\ref.docx'
doc = docx.Document(ref_path)

print("=== SECTIONS & MARGINS ===")
for i, sec in enumerate(doc.sections):
    print(f"Section {i}:")
    print(f"  Page Width: {sec.page_width.inches if sec.page_width else None} in ({sec.page_width.pt if sec.page_width else None} pt)")
    print(f"  Page Height: {sec.page_height.inches if sec.page_height else None} in ({sec.page_height.pt if sec.page_height else None} pt)")
    print(f"  Orientation: {sec.orientation}")
    print(f"  Top Margin: {sec.top_margin.inches if sec.top_margin else None} in")
    print(f"  Bottom Margin: {sec.bottom_margin.inches if sec.bottom_margin else None} in")
    print(f"  Left Margin: {sec.left_margin.inches if sec.left_margin else None} in")
    print(f"  Right Margin: {sec.right_margin.inches if sec.right_margin else None} in")
    print(f"  Header Distance: {sec.header_distance.inches if sec.header_distance else None} in")
    print(f"  Footer Distance: {sec.footer_distance.inches if sec.footer_distance else None} in")
    print(f"  Different First Page: {sec.different_first_page_header_footer}")
    
    sectPr = sec._sectPr
    pgBorders = sectPr.find(qn('w:pgBorders'))
    if pgBorders is not None:
        print(f"  Page Borders XML: {pgBorders.xml}")
    else:
        print("  Page Borders: None")

print("\n=== STYLES ===")
for style in doc.styles:
    if style.type == docx.enum.style.WD_STYLE_TYPE.PARAGRAPH:
        if style.name in ['Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'Title', 'Subtitle', 'Header', 'Footer']:
            font = style.font
            pf = style.paragraph_format
            color_str = str(font.color.rgb) if font.color and font.color.rgb else 'None'
            print(f"Style '{style.name}': font={font.name}, size={font.size.pt if font.size else None} pt, bold={font.bold}, color={color_str}, align={pf.alignment}, ls={pf.line_spacing}, sb={pf.space_before.pt if pf.space_before else None}, sa={pf.space_after.pt if pf.space_after else None}")

print("\n=== HEADERS & FOOTERS ===")
for i, sec in enumerate(doc.sections):
    h = sec.header
    f = sec.footer
    h_text = ' '.join(p.text.strip() for p in h.paragraphs if p.text.strip())
    f_text = ' '.join(p.text.strip() for p in f.paragraphs if p.text.strip())
    print(f"Section {i} Header: '{h_text}'")
    print(f"Section {i} Footer: '{f_text}'")
    if sec.different_first_page_header_footer:
        fh = sec.first_page_header
        ff = sec.first_page_footer
        print(f"Section {i} First Page Header: '{' '.join(p.text.strip() for p in fh.paragraphs if p.text.strip())}'")
        print(f"Section {i} First Page Footer: '{' '.join(p.text.strip() for p in ff.paragraphs if p.text.strip())}'")

print("\n=== SAMPLE PARAGRAPHS (First 40) ===")
for i, p in enumerate(doc.paragraphs[:40]):
    t = p.text.strip()
    if not t and len(p.runs) == 0:
        continue
    runs_info = []
    for r in p.runs[:4]:
        c = str(r.font.color.rgb) if r.font.color and r.font.color.rgb else 'None'
        runs_info.append((r.text[:20], r.font.name, r.font.size.pt if r.font.size else None, r.font.bold, c))
    pf = p.paragraph_format
    sb = pf.space_before.pt if pf.space_before else None
    sa = pf.space_after.pt if pf.space_after else None
    ls = pf.line_spacing
    al = p.alignment
    has_img = '<w:drawing' in p._element.xml
    print(f"P[{i:2d}] style='{p.style.name}' al={al} sb={sb} sa={sa} ls={ls} img={has_img} | text='{t[:40]}' | runs={runs_info}")

print(f"\nTotal paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

print("\n=== SAMPLE TABLES ===")
for t_idx, tbl in enumerate(doc.tables[:5]):
    print(f"Table {t_idx}: rows={len(tbl.rows)}, cols={len(tbl.columns)}, style={tbl.style.name}")
    row0 = tbl.rows[0]
    headers = []
    for cell in row0.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.find(qn('w:shd'))
        fill = shd.attrib.get(qn('w:fill')) if shd is not None else None
        p_runs = []
        for p in cell.paragraphs:
            for r in p.runs:
                c = str(r.font.color.rgb) if r.font.color and r.font.color.rgb else 'None'
                p_runs.append((r.text.strip(), r.font.name, r.font.size.pt if r.font.size else None, r.font.bold, c))
        headers.append((cell.text.strip(), fill, p_runs))
    print(f"  Header Row: {headers}")
    if len(tbl.rows) > 1:
        row1 = tbl.rows[1]
        r1_cells = []
        for cell in row1.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = tcPr.find(qn('w:shd'))
            fill = shd.attrib.get(qn('w:fill')) if shd is not None else None
            p_runs = []
            for p in cell.paragraphs:
                for r in p.runs:
                    c = str(r.font.color.rgb) if r.font.color and r.font.color.rgb else 'None'
                    p_runs.append((r.text.strip(), r.font.name, r.font.size.pt if r.font.size else None, r.font.bold, c))
            r1_cells.append((cell.text.strip()[:20], fill, p_runs))
        print(f"  Row 1: {r1_cells}")
    tblPr = tbl._tbl.tblPr
    tblBorders = tblPr.find(qn('w:tblBorders'))
    if tblBorders is not None:
        print(f"  Table Borders XML: {tblBorders.xml}")
